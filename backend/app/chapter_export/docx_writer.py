"""DOCX 导出 writer - 将导出计划内容生成 Word 文档。"""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from docx import Document

from app.chapter_export.service import chinese_number

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _add_title(document: Document, text: str, level: int) -> None:
    heading = document.add_heading("", level=level)
    heading.add_run(text)


def _add_body(document: Document, content: str) -> None:
    for line in content.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        document.add_paragraph(line)


async def write_docx_export(
    output_path: Path,
    *,
    mode: str,
    chapters: list[dict[str, Any]],
    volume_by_chapter: dict[str, dict[str, Any]],
    load_batch: Any,
) -> int:
    """构建并保存 DOCX 导出文件，返回写入章节数。

    load_batch 是异步回调：接收一批章节 id，返回章节实体列表，
    由调用方提供分批读取与取消检查逻辑。
    """
    document = Document()
    written_count = 0
    last_volume_id: str | None = None

    for offset in range(0, len(chapters), 20):
        batch = chapters[offset : offset + 20]
        ids = [item.get("id") for item in batch if isinstance(item.get("id"), str)]
        loaded_by_id = {chapter.id: chapter for chapter in await load_batch(ids)}

        for item in batch:
            chapter_id = item.get("id")
            if not isinstance(chapter_id, str):
                raise ValueError("导出任务章节数据无效")
            chapter = loaded_by_id.get(chapter_id)
            if chapter is None:
                raise ValueError("导出章节已被删除，请重新发起导出")
            title = item.get("title") if isinstance(item.get("title"), str) else chapter.title

            volume = volume_by_chapter.get(chapter_id)
            if mode == "volumes":
                if volume is None:
                    raise ValueError("导出任务卷数据无效")
                if volume.get("id") != last_volume_id:
                    order = volume.get("order")
                    volume_title = volume.get("title")
                    if not isinstance(order, int) or not isinstance(volume_title, str):
                        raise ValueError("导出任务卷数据无效")
                    _add_title(document, f"第{chinese_number(order)}卷 {volume_title}", 1)
                    last_volume_id = volume.get("id")
                _add_title(document, title, 2)
            else:
                _add_title(document, title, 1)

            _add_body(document, chapter.content or "")
            written_count += 1

    document.core_properties.title = "章节导出"
    await asyncio.to_thread(document.save, str(output_path))
    return written_count
