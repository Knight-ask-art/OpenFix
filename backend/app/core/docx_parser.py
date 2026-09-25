# -*- coding: utf-8 -*-
"""DOCX 解析 - 从 Word 文档提取文本并复用现有章节切分逻辑。"""

import io

from docx import Document
from docx.text.paragraph import Paragraph

from app.core.txt_parser import (
    ParseResult,
    ParsedChapter,
    ParsedVolume,
    _count_words,
    parse_txt_content,
)

_DEFAULT_VOLUME_TITLE = "第一卷"
_MAX_TITLE_LENGTH = 50


def _paragraph_style_name(paragraph: Paragraph) -> str:
    try:
        style = paragraph.style
    except Exception:
        return ""
    return (style.name or "") if style is not None else ""


def _heading_level(style_name: str) -> int | None:
    """返回 Heading 样式层级；非标题样式返回 None。

    兼容英文（Heading 1）与中文（标题 1）样式名。
    """
    lowered = style_name.lower()
    if lowered.startswith("heading") or lowered.startswith("标题"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        return int(digits) if digits else 1
    return None


def _collect_paragraphs(document: Document) -> list[tuple[int | None, str]]:
    collected: list[tuple[int | None, str]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        collected.append((_heading_level(_paragraph_style_name(paragraph)), text))
    return collected


def extract_docx_text(content: bytes) -> str:
    """提取 DOCX 全部段落文本（换行分隔），用于手动字数切分。"""
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("DOCX 文件无法读取，请确认文件没有损坏或加密") from exc
    return "\n".join(text for _, text in _collect_paragraphs(document))


def parse_docx_content(content: bytes) -> ParseResult:
    """解析 DOCX 文件内容，返回统一的卷章节结构。

    优先按文档内置的 Heading 样式切分：Heading 1 作为卷，Heading 2 及以下
    作为章；文档没有 Heading 样式时，将全文交给现有 TXT 章节规则识别。
    """
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("DOCX 文件无法读取，请确认文件没有损坏或加密") from exc

    paragraphs = _collect_paragraphs(document)
    if not paragraphs:
        return ParseResult(detected_encoding="docx")

    if any(level is not None for level, _ in paragraphs):
        return _parse_with_headings(paragraphs)
    return _parse_without_headings(paragraphs)


def _parse_with_headings(paragraphs: list[tuple[int | None, str]]) -> ParseResult:
    volumes: list[ParsedVolume] = []
    current_volume: ParsedVolume | None = None
    current_chapter: ParsedChapter | None = None
    body_buffer: list[str] = []

    def flush_chapter() -> None:
        nonlocal current_chapter, body_buffer
        if current_chapter is not None:
            content = "\n".join(body_buffer).strip()
            current_chapter.content = content
            current_chapter.word_count = _count_words(content)
        body_buffer = []

    def ensure_volume(title: str) -> ParsedVolume:
        nonlocal current_volume
        volume = ParsedVolume(title=title[:_MAX_TITLE_LENGTH])
        volumes.append(volume)
        current_volume = volume
        return volume

    for level, text in paragraphs:
        if level is None:
            body_buffer.append(text)
            continue

        flush_chapter()
        current_chapter = None
        if level == 1:
            ensure_volume(text)
            continue

        if current_volume is None:
            ensure_volume(_DEFAULT_VOLUME_TITLE)
        current_chapter = ParsedChapter(title=text[:_MAX_TITLE_LENGTH], content="", word_count=0)
        current_volume.chapters.append(current_chapter)

    flush_chapter()

    chapters = [chapter for volume in volumes for chapter in volume.chapters]
    if not chapters:
        # 只有卷标题、没有任何章节内容时退回文本规则识别。
        return _parse_without_headings(paragraphs)
    return _finalize(volumes, "docx")


def _parse_without_headings(paragraphs: list[tuple[int | None, str]]) -> ParseResult:
    text = "\n".join(text for _, text in paragraphs)
    if not text.strip():
        return ParseResult(detected_encoding="docx")
    result = parse_txt_content(text.encode("utf-8"))
    result.detected_encoding = "docx"
    return result


def _finalize(volumes: list[ParsedVolume], encoding: str) -> ParseResult:
    chapters = [chapter for volume in volumes for chapter in volume.chapters]
    return ParseResult(
        volumes=volumes,
        total_word_count=sum(chapter.word_count for chapter in chapters),
        chapter_count=len(chapters),
        detected_encoding=encoding,
    )
