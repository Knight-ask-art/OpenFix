# -*- coding: utf-8 -*-
"""章节导出 API 测试。"""

import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from httpx import AsyncClient
from urllib.parse import unquote

from app.background.events.publisher import BackgroundEventPublisher
from app.background.jobs import service as background_service
from app.background.jobs.definitions.chapter_export import (
    CHAPTER_EXPORT_JOB,
    cleanup_chapter_export,
)
from app.background.runtime.context import JobContext
from app.background.runtime.dispatcher import dispatch_job
from app.chapter_export import service as chapter_export_service
from app.background.jobs.models import BackgroundJob
from app.api.routers import chapter_exports as chapter_exports_router
from app.storage.repos import chapter_repo


def test_chinese_volume_numbers() -> None:
    assert chapter_export_service.chinese_number(1) == "一"
    assert chapter_export_service.chinese_number(10) == "十"
    assert chapter_export_service.chinese_number(11) == "十一"
    assert chapter_export_service.chinese_number(21) == "二十一"
    assert chapter_export_service.chinese_number(101) == "一百零一"


def test_expired_export_is_not_downloadable(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job = BackgroundJob(
        id="expired-export",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="succeeded",
        payload_json='{"filename":"测试.txt"}',
        result_json='{"expires_at":"2020-01-01T00:00:00+00:00"}',
    )
    _part_path, output_path = chapter_export_service.export_file_paths(job.id)
    output_path.write_text("expired", encoding="utf-8")

    assert not chapter_export_service.is_export_download_available(job)


async def _create_project(client: AsyncClient, title: str = "测试小说") -> tuple[str, str]:
    response = await client.post("/api/v1/projects", data={"title": title})
    assert response.status_code == 201
    project_id = response.json()["id"]
    volumes = (await client.get(f"/api/v1/projects/{project_id}/volumes")).json()
    return project_id, volumes[0]["id"]


async def _create_chapter(
    client: AsyncClient,
    project_id: str,
    volume_id: str,
    title: str,
    content: str,
    word_count: int,
) -> dict:
    response = await client.post(
        f"/api/v1/projects/{project_id}/chapters",
        json={
            "volume_id": volume_id,
            "title": title,
            "content": content,
            "word_count": word_count,
        },
    )
    assert response.status_code == 201
    return response.json()


@pytest.mark.asyncio
async def test_create_full_volume_export_uses_volume_filename_and_snapshot_selection(
    client: AsyncClient,
) -> None:
    project_id, volume_id = await _create_project(client)
    first = await _create_chapter(client, project_id, volume_id, "第一章", "第一章正文\r\n第二行", 5)
    second = await _create_chapter(client, project_id, volume_id, "第二章", "第二章正文", 5)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "pending"
    assert data["chapter_count"] == 2
    assert data["word_count"] == 10
    assert data["filename"] == "测试小说-全本-2026-07-28.txt"
    assert data["chapter_ids"] == [first["id"], second["id"]]


@pytest.mark.asyncio
async def test_export_creation_does_not_load_chapter_bodies(client: AsyncClient, monkeypatch) -> None:
    project_id, volume_id = await _create_project(client)
    chapter = await _create_chapter(client, project_id, volume_id, "第一章", "正文", 2)

    async def reject_full_chapter_load(*_args, **_kwargs):
        raise AssertionError("导出创建阶段不应读取完整章节正文")

    monkeypatch.setattr(chapter_repo, "list_by_project", reject_full_chapter_load)
    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [chapter["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201


@pytest.mark.asyncio
async def test_only_cancel_endpoint_preempts_running_export(client: AsyncClient, monkeypatch) -> None:
    class Supervisor:
        def __init__(self) -> None:
            self.cancelled_job_ids: list[str] = []

        def create_event_publisher(self) -> BackgroundEventPublisher:
            return BackgroundEventPublisher(None)

        def cancel_running_chapter_export(self, job_id: str) -> bool:
            self.cancelled_job_ids.append(job_id)
            return False

    supervisor = Supervisor()
    monkeypatch.setattr(chapter_exports_router, "get_background_supervisor", lambda: supervisor)
    project_id, volume_id = await _create_project(client)
    chapter = await _create_chapter(client, project_id, volume_id, "第一章", "正文", 2)

    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [chapter["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )
    assert created.status_code == 201
    assert supervisor.cancelled_job_ids == []

    cancelled = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports/{created.json()['id']}/cancel"
    )
    assert cancelled.status_code == 200
    assert supervisor.cancelled_job_ids == [created.json()["id"]]


@pytest.mark.asyncio
async def test_create_fragment_export_uses_chapter_filename(client: AsyncClient) -> None:
    project_id, volume_id = await _create_project(client)
    selected = await _create_chapter(client, project_id, volume_id, "第一章", "正文", 2)
    await _create_chapter(client, project_id, volume_id, "第二章", "正文", 2)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [selected["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == "测试小说-1个章节-2026-07-28.txt"


@pytest.mark.asyncio
async def test_manually_selected_complete_volume_still_uses_chapter_format(
    client: AsyncClient,
) -> None:
    project_id, volume_id = await _create_project(client)
    first = await _create_chapter(client, project_id, volume_id, "第一章", "正文", 2)
    second = await _create_chapter(client, project_id, volume_id, "第二章", "正文", 2)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [first["id"], second["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == "测试小说-2个章节-2026-07-28.txt"


@pytest.mark.asyncio
async def test_create_export_rejects_empty_selection(client: AsyncClient) -> None:
    project_id, _volume_id = await _create_project(client)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 400
    assert "章节" in response.json()["detail"]


@pytest.mark.asyncio
async def test_export_task_writes_full_volume_txt_and_serves_download(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    first = await _create_chapter(client, project_id, volume_id, "第一章", "第一章正文\r\n第二行", 5)
    second = await _create_chapter(client, project_id, volume_id, "第二章", "第二章正文", 5)
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )
    assert created.status_code == 201
    job = await background_service.get_job(session, created.json()["id"])
    assert job is not None
    job.status = "running"
    await session.commit()

    context = JobContext(session=session, job=job, publisher=BackgroundEventPublisher(None))
    result = await dispatch_job(context)
    await background_service.mark_succeeded(session, context.publisher, context.job, result=result)
    await session.commit()

    # 真实 dispatch + mark_succeeded 之后立即清理，成品必须仍然存在且可下载。
    assert await chapter_export_service.cleanup_chapter_export_files(session) == 0
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "txt")
    assert output_path.is_file()

    status_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}"
    )
    assert status_response.status_code == 200
    assert status_response.json()["current"] == 2
    assert status_response.json()["total"] == 2
    assert status_response.json()["download_url"]

    download_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}/download"
    )
    assert download_response.status_code == 200
    assert "attachment" in download_response.headers["content-disposition"]
    assert "测试小说-全本-2026-07-28.txt" in unquote(
        download_response.headers["content-disposition"]
    )
    assert download_response.content.decode("utf-8-sig") == (
        "第一卷 第一卷\n"
        "第一章\n第一章正文\n第二行\n\n第二章\n第二章正文"
    )
    assert result == {
        "filename": "测试小说-全本-2026-07-28.txt",
        "volume_count": 1,
        "chapter_count": 2,
        "word_count": 10,
        "expires_at": result["expires_at"],
    }
    assert [first["id"], second["id"]] == created.json()["chapter_ids"]


@pytest.mark.asyncio
@pytest.mark.parametrize("export_format", ["txt", "docx", "markdown"])
async def test_cancelled_export_removes_partial_file(
    client: AsyncClient,
    monkeypatch,
    tmp_path,
    export_format: str,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "第一章", "正文", 2)
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
            "format": export_format,
        },
    )
    assert created.status_code == 201
    assert created.json()["format"] == export_format
    job_id = created.json()["id"]
    part_path, output_path = chapter_export_service.export_file_paths(job_id, export_format)
    part_path.write_text("partial", encoding="utf-8")
    output_path.write_text("complete", encoding="utf-8")

    cancelled = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports/{job_id}/cancel"
    )

    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    # 立即删除接缝：part 与该格式成品必须同时消失，不能等到周期清理。
    assert not part_path.exists()
    assert not output_path.exists()
    assert [path.name for path in tmp_path.iterdir()] == []


@pytest.mark.asyncio
async def test_cleanup_keeps_output_while_export_is_still_running(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job = BackgroundJob(
        id="running-export",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="running",
        payload_json="{}",
    )
    session.add(job)
    await session.commit()
    _part_path, output_path = chapter_export_service.export_file_paths(job.id)
    output_path.write_text("finished but not committed", encoding="utf-8")

    assert await chapter_export_service.cleanup_chapter_export_files(session) == 0
    assert output_path.exists()


@pytest.mark.asyncio
async def test_export_task_writes_full_volume_docx_and_serves_download(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    import io

    from docx import Document

    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "第一章", "第一章正文\r\n第二行", 5)
    await _create_chapter(client, project_id, volume_id, "第二章", "第二章正文", 5)
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
            "format": "docx",
        },
    )
    assert created.status_code == 201
    assert created.json()["filename"] == "测试小说-全本-2026-07-28.docx"
    assert created.json()["format"] == "docx"

    job = await background_service.get_job(session, created.json()["id"])
    assert job is not None
    job.status = "running"
    await session.commit()

    context = JobContext(session=session, job=job, publisher=BackgroundEventPublisher(None))
    result = await dispatch_job(context)
    await background_service.mark_succeeded(session, context.publisher, context.job, result=result)
    await session.commit()

    # 真实 dispatch + mark_succeeded 之后立即清理，DOCX 成品必须仍然存在且可下载。
    assert await chapter_export_service.cleanup_chapter_export_files(session) == 0
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "docx")
    assert output_path.is_file()

    status_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}"
    )
    assert status_response.status_code == 200
    assert status_response.json()["download_url"]

    download_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}/download"
    )
    assert download_response.status_code == 200
    assert (
        download_response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "测试小说-全本-2026-07-28.docx" in unquote(
        download_response.headers["content-disposition"]
    )

    document = Document(io.BytesIO(download_response.content))
    headings = [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading")]
    body_texts = [p.text for p in document.paragraphs if not p.style.name.startswith("Heading") and p.text.strip()]
    assert ("Heading 1", "第一卷 第一卷") in headings
    assert ("Heading 2", "第一章") in headings
    assert ("Heading 2", "第二章") in headings
    assert "第二行" in body_texts
    assert "第二章正文" in body_texts


@pytest.mark.asyncio
async def test_create_docx_export_rejects_invalid_format(client: AsyncClient) -> None:
    project_id, _volume_id = await _create_project(client)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
            "format": "pdf",
        },
    )

    assert response.status_code == 422


def _write_export_artifacts(directory: Path, job_id: str, export_format: str) -> dict[str, Path]:
    """写入 part、匹配格式成品和其它已识别后缀成品，覆盖改名窗口的真实组合。"""
    other_suffix = ".docx" if export_format == "txt" else ".txt"
    suffix = ".md" if export_format == "markdown" else f".{export_format}"
    paths = {
        "part": directory / f"{chapter_export_service.EXPORT_FILE_PREFIX}{job_id}.part",
        "matching": directory / f"{chapter_export_service.EXPORT_FILE_PREFIX}{job_id}{suffix}",
        "mismatched": directory / f"{chapter_export_service.EXPORT_FILE_PREFIX}{job_id}{other_suffix}",
    }
    for path in paths.values():
        path.write_text("synthetic", encoding="utf-8")
    return paths


async def _add_export_job(
    session,
    job_id: str,
    status: str,
    payload: dict,
    result: dict | None = None,
    *,
    job_type: str = chapter_export_service.EXPORT_JOB_TYPE,
) -> None:
    session.add(
        BackgroundJob(
            id=job_id,
            type=job_type,
            status=status,
            payload_json=json.dumps(payload),
            result_json=json.dumps(result) if result is not None else None,
        )
    )
    await session.commit()


_HOOK_STATUS = {
    "on_failed": "failed",
    "on_timeout": "timeout",
    "on_cancelled": "cancelled",
}


@pytest.mark.asyncio
@pytest.mark.parametrize("hook_name", ["on_failed", "on_timeout", "on_cancelled"])
@pytest.mark.parametrize(
    ("payload_format", "expected_format"),
    [("txt", "txt"), ("docx", "docx"), ("markdown", "markdown"), (None, "txt")],
)
async def test_terminal_hook_deletes_immediate_export_artifacts(
    session,
    monkeypatch,
    tmp_path,
    hook_name: str,
    payload_format: str | None,
    expected_format: str,
) -> None:
    """已注册的终态钩子必须按任务格式立即删除 part 与该格式成品。

    失败与超时由 worker/watchdog 驱动，这里直接调用章节目录上真实注册的钩子来覆盖
    该接缝；取消接缝已由上面的真实 API 取消测试覆盖。不 mock 删除函数，也不调用周期清理。
    """
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job_id = f"terminal-{hook_name}-{expected_format}-{payload_format}"
    payload: dict = {"filename": f"导出.{expected_format}"}
    if payload_format is not None:
        payload["format"] = payload_format
    status = _HOOK_STATUS[hook_name]
    await _add_export_job(session, job_id, status, payload)
    paths = _write_export_artifacts(tmp_path, job_id, expected_format)
    unrelated = tmp_path / f"{chapter_export_service.EXPORT_FILE_PREFIX}unrelated-job.part"
    unrelated.write_text("synthetic", encoding="utf-8")

    job = await background_service.get_job(session, job_id)
    assert job is not None
    context = JobContext(
        session=session,
        job=job,
        publisher=BackgroundEventPublisher(None),
        definition=CHAPTER_EXPORT_JOB,
    )
    hook = getattr(CHAPTER_EXPORT_JOB, hook_name)
    assert hook is cleanup_chapter_export
    await hook(context, "终态接缝测试")

    assert not paths["part"].exists()
    assert not paths["matching"].exists()
    # 非本格式后缀不是该任务的成品，只有周期清理负责；它不应被立即删除接缝误伤。
    assert paths["mismatched"].exists()
    assert unrelated.exists()
    assert job.status == status


@pytest.mark.asyncio
@pytest.mark.parametrize("export_format", ["txt", "docx", "markdown"])
@pytest.mark.parametrize(
    ("case", "status", "result", "keep"),
    [
        ("pending", "pending", None, (True, True, False)),
        ("running", "running", None, (True, True, False)),
        ("cancel-requested", "cancel_requested", None, (True, True, False)),
        (
            "succeeded-fresh",
            "succeeded",
            {"expires_at": "2999-01-01T00:00:00+00:00"},
            (False, True, False),
        ),
        (
            "succeeded-expired",
            "succeeded",
            {"expires_at": "2000-01-01T00:00:00+00:00"},
            (False, False, False),
        ),
        ("succeeded-invalid-expiry", "succeeded", {"expires_at": "not-a-timestamp"}, (False, False, False)),
        ("succeeded-missing-expiry", "succeeded", {}, (False, False, False)),
        ("failed", "failed", None, (False, False, False)),
        ("cancelled", "cancelled", None, (False, False, False)),
        ("timeout", "timeout", None, (False, False, False)),
        ("skipped", "skipped", None, (False, False, False)),
    ],
)
async def test_cleanup_applies_export_artifact_lifecycle_matrix(
    session,
    monkeypatch,
    tmp_path,
    export_format: str,
    case: str,
    status: str,
    result: dict | None,
    keep: tuple[bool, bool, bool],
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job_id = f"lifecycle-{case}-{export_format}"
    await _add_export_job(
        session,
        job_id,
        status,
        {"filename": f"导出.{export_format}", "format": export_format},
        result,
    )
    paths = _write_export_artifacts(tmp_path, job_id, export_format)

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == keep.count(False)
    assert {name: path.exists() for name, path in paths.items()} == {
        "part": keep[0],
        "matching": keep[1],
        "mismatched": keep[2],
    }


@pytest.mark.asyncio
async def test_cleanup_removes_orphan_and_foreign_type_but_keeps_unknown_names(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    await _add_export_job(session, "foreign-job", "running", {"format": "txt"}, job_type="backup")
    prefix = chapter_export_service.EXPORT_FILE_PREFIX
    removed_paths = [
        tmp_path / f"{prefix}missing-job.part",
        tmp_path / f"{prefix}missing-job.txt",
        tmp_path / f"{prefix}missing-job.docx",
        tmp_path / f"{prefix}missing-job.md",
        tmp_path / f"{prefix}foreign-job.part",
        tmp_path / f"{prefix}foreign-job.txt",
        tmp_path / f"{prefix}foreign-job.md",
    ]
    kept_paths = [
        tmp_path / f"{prefix}missing-job",
        tmp_path / f"{prefix}missing-job.zip",
        tmp_path / f"{prefix}missing-job.doc",
        tmp_path / f"{prefix}.txt",
        tmp_path / f"{prefix}.md",
        tmp_path / f"{prefix}missing-job.txt.bak",
        tmp_path / "notes.txt",
        tmp_path / "missing-job.txt",
    ]
    for path in [*removed_paths, *kept_paths]:
        path.write_text("synthetic", encoding="utf-8")

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == len(removed_paths)
    assert all(not path.exists() for path in removed_paths)
    assert all(path.exists() for path in kept_paths)


@pytest.mark.asyncio
async def test_cleanup_treats_missing_format_as_legacy_txt(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    await _add_export_job(session, "legacy-running", "running", {"filename": "导出.txt"})
    await _add_export_job(
        session,
        "legacy-succeeded",
        "succeeded",
        {"filename": "导出.txt"},
        {"expires_at": "2999-01-01T00:00:00+00:00"},
    )
    running_paths = _write_export_artifacts(tmp_path, "legacy-running", "txt")
    succeeded_paths = _write_export_artifacts(tmp_path, "legacy-succeeded", "txt")
    markdown_paths = [tmp_path / f"chapter-export-{job_id}.md" for job_id in ("legacy-running", "legacy-succeeded")]
    for path in markdown_paths:
        path.write_text("synthetic", encoding="utf-8")

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == 5
    assert all(not path.exists() for path in markdown_paths)
    assert running_paths["part"].exists()
    assert running_paths["matching"].exists()
    assert not running_paths["mismatched"].exists()
    assert not succeeded_paths["part"].exists()
    assert succeeded_paths["matching"].exists()
    assert not succeeded_paths["mismatched"].exists()


@pytest.mark.asyncio
@pytest.mark.parametrize("invalid_format", ["pdf", None, 7])
@pytest.mark.parametrize("export_format", ["txt", "docx", "markdown"])
async def test_cleanup_does_not_preserve_completed_artifact_for_invalid_format(
    session,
    monkeypatch,
    tmp_path,
    invalid_format,
    export_format: str,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job_id = f"invalid-format-{invalid_format}"
    await _add_export_job(
        session,
        job_id,
        "succeeded",
        {"filename": "导出.txt", "format": invalid_format},
        {"expires_at": "2999-01-01T00:00:00+00:00"},
    )
    paths = _write_export_artifacts(tmp_path, job_id, export_format)

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == 3
    assert not any(path.exists() for path in paths.values())


@pytest.mark.asyncio
@pytest.mark.parametrize("export_format", ["txt", "docx", "markdown"])
async def test_cleanup_keeps_in_flight_part_for_active_job_with_invalid_format(
    session,
    monkeypatch,
    tmp_path,
    export_format: str,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job_id = "invalid-format-running"
    await _add_export_job(
        session,
        job_id,
        "running",
        {"filename": "导出.pdf", "format": "pdf"},
    )
    paths = _write_export_artifacts(tmp_path, job_id, export_format)

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == 2
    assert paths["part"].exists()
    assert not paths["matching"].exists()
    assert not paths["mismatched"].exists()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("case", "expires_at", "keep_matching"),
    [
        ("naive-future", "2999-01-01T00:00:00", True),
        ("offset-future", "2999-01-01T08:00:00+08:00", True),
        ("naive-past", "2000-01-01T00:00:00", False),
    ],
)
@pytest.mark.parametrize("export_format", ["txt", "docx", "markdown"])
async def test_cleanup_interprets_naive_expiry_as_utc(
    session,
    monkeypatch,
    tmp_path,
    case: str,
    expires_at: str,
    keep_matching: bool,
    export_format: str,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job_id = f"naive-expiry-{case}"
    await _add_export_job(
        session,
        job_id,
        "succeeded",
        {"format": export_format},
        {"expires_at": expires_at},
    )
    paths = _write_export_artifacts(tmp_path, job_id, export_format)

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == (2 if keep_matching else 3)
    assert paths["matching"].exists() == keep_matching
    assert not paths["part"].exists()


@pytest.mark.asyncio
@pytest.mark.parametrize("export_format", ["txt", "docx", "markdown"])
async def test_cleanup_expiry_boundary_requires_strictly_future_expiry(
    session,
    monkeypatch,
    tmp_path,
    export_format: str,
) -> None:
    frozen_now = datetime(2030, 1, 1, tzinfo=UTC)

    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return frozen_now

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(chapter_export_service, "datetime", FrozenDateTime)
    await _add_export_job(
        session,
        "boundary-equal",
        "succeeded",
        {"format": export_format},
        {"expires_at": "2030-01-01T00:00:00+00:00"},
    )
    await _add_export_job(
        session,
        "boundary-after",
        "succeeded",
        {"format": export_format},
        {"expires_at": "2030-01-01T00:00:01+00:00"},
    )
    equal_paths = _write_export_artifacts(tmp_path, "boundary-equal", export_format)
    after_paths = _write_export_artifacts(tmp_path, "boundary-after", export_format)

    removed = await chapter_export_service.cleanup_chapter_export_files(session)

    assert removed == 5
    assert not equal_paths["part"].exists()
    assert not equal_paths["matching"].exists()
    assert not equal_paths["mismatched"].exists()
    assert not after_paths["part"].exists()
    assert not after_paths["mismatched"].exists()
    assert after_paths["matching"].exists()
