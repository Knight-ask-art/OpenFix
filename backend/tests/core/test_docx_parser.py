# -*- coding: utf-8 -*-
"""DOCX 解析测试。"""

import io

import pytest
from docx import Document

from app.core.docx_parser import extract_docx_text, parse_docx_content
from app.core.project_import import parse_project_import


def _build_docx(paragraphs: list[tuple[str | None, str]]) -> bytes:
    """构造 DOCX 字节流。paragraphs 为 (样式名或 None, 文本) 列表。"""
    document = Document()
    for style, text in paragraphs:
        paragraph = document.add_paragraph()
        if style:
            paragraph.style = document.styles[style]
        paragraph.add_run(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_docx_with_heading_style() -> None:
    content = _build_docx(
        [
            ("Heading 1", "第一卷"),
            ("Heading 2", "第一章 起点"),
            (None, "他睁开眼，发现自己回到了十年前。" * 3),
            ("Heading 2", "第二章 风起"),
            (None, "江湖传言，剑冢重现于世。" * 3),
        ]
    )

    result = parse_docx_content(content)

    assert result.detected_encoding == "docx"
    assert result.chapter_count == 2
    assert [volume.title for volume in result.volumes] == ["第一卷"]
    assert [chapter.title for chapter in result.volumes[0].chapters] == [
        "第一章 起点",
        "第二章 风起",
    ]
    assert "剑冢重现" in result.volumes[0].chapters[1].content
    assert result.total_word_count > 0


def test_parse_docx_without_headings_uses_text_rules() -> None:
    body_1 = "夜色笼罩着小镇，远处传来犬吠。" * 40
    body_2 = "清晨的薄雾还未散去，商队的驼铃已经响起。" * 40
    content = _build_docx(
        [
            (None, "第一章 开始"),
            (None, body_1),
            (None, "第二章 发展"),
            (None, body_2),
        ]
    )

    result = parse_docx_content(content)

    assert result.chapter_count >= 2
    titles = [chapter.title for volume in result.volumes for chapter in volume.chapters]
    assert any("第一章" in title for title in titles)
    assert any("第二章" in title for title in titles)
    assert result.detected_encoding == "docx"


def test_parse_docx_single_body_paragraph() -> None:
    content = _build_docx([(None, "只有一段没有标题的文字。")])

    result = parse_docx_content(content)

    assert result.chapter_count == 1
    assert result.volumes[0].chapters[0].title in {"只有一段没有标题的文字。", "正文"}


def test_parse_docx_invalid_file_raises_value_error() -> None:
    with pytest.raises(ValueError, match="DOCX"):
        parse_docx_content(b"not a docx file")


def test_extract_docx_text_joins_paragraphs() -> None:
    content = _build_docx(
        [
            (None, "第一段。"),
            (None, "第二段。"),
        ]
    )

    assert extract_docx_text(content) == "第一段。\n第二段。"


def test_parse_project_import_supports_docx() -> None:
    content = _build_docx(
        [
            ("Heading 1", "第一卷"),
            ("Heading 2", "第一章"),
            (None, "正文内容若干。" * 10),
        ]
    )

    result = parse_project_import("novel.docx", content)

    assert result.chapter_count == 1
    assert result.volumes[0].title == "第一卷"


def test_parse_project_import_docx_manual_split() -> None:
    content = _build_docx(
        [
            (None, "这是一个用于手动切分的段落。" * 5),
            (None, "第二段内容同样用于切分测试。" * 5),
        ]
    )

    result = parse_project_import("novel.docx", content, split_mode="manual", chunk_size=50)

    assert result.chapter_count >= 2
    assert result.detected_encoding == "utf-8"
